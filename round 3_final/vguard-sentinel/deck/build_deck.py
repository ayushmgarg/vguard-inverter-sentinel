"""Builds VGuard-Sentinel-Finale-Deck.pptx in the V-Guard Big Idea Tech palette (black / gold #FDC300 / orange #F39200 / white)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os
GOLD, ORANGE, BLACK, INK, GREY, LIGHT, WHITE = [RGBColor.from_string(c) for c in ("FDC300", "F39200", "000104", "221E1F", "6C757D", "F3EFEF", "FFFFFF")]
A = "assets"; HEAD = "Arial"; BODY = "Arial"
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = prs.slide_width, prs.slide_height
n_slides = [0]

def rect(slide, x, y, w, h, fill, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h); s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else: s.line.color.rgb = line
    s.shadow.inherit = False; return s

def text(slide, x, y, w, h, txt, size=18, bold=False, color=INK, font=BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
    lines = txt if isinstance(txt, list) else [txt]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
        r = p.add_run(); r.text = ln; f = r.font; f.size = Pt(size); f.bold = bold; f.italic = italic; f.color.rgb = color; f.name = font
        p.space_after = Pt(4)
    return tb

def bullets(slide, x, y, w, h, items, size=16, color=INK, gap=6):
    tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r0 = p.add_run(); r0.text = "▪  "; r0.font.color.rgb = ORANGE; r0.font.size = Pt(size); r0.font.name = BODY
        if it.startswith("**") and "**" in it[2:]:
            head, rest = it[2:].split("**", 1)
            r1 = p.add_run(); r1.text = head; r1.font.bold = True; r1.font.size = Pt(size); r1.font.color.rgb = color; r1.font.name = BODY
            r2 = p.add_run(); r2.text = rest; r2.font.size = Pt(size); r2.font.color.rgb = color; r2.font.name = BODY
        else:
            r1 = p.add_run(); r1.text = it; r1.font.size = Pt(size); r1.font.color.rgb = color; r1.font.name = BODY
        p.space_after = Pt(gap)
    return tb

def chrome(slide, title, kicker=None):
    n_slides[0] += 1
    rect(slide, 0, 0, W, Inches(0.12), GOLD)
    rect(slide, Inches(0.6), Inches(0.42), Inches(0.12), Inches(0.7), ORANGE)
    text(slide, Inches(0.85), Inches(0.32), Inches(11.5), Inches(0.9), title, size=30, bold=True, font=HEAD, color=BLACK)
    if kicker: text(slide, Inches(0.87), Inches(1.02), Inches(11.5), Inches(0.5), kicker, size=14, color=GREY, italic=True)
    rect(slide, 0, H - Inches(0.42), W, Inches(0.42), BLACK)
    text(slide, Inches(0.6), H - Inches(0.40), Inches(9), Inches(0.38), "V-Guard Sentinel  ·  Big Idea Tech 2026, Track 4  ·  Team Codey Tingle (TI3271)", size=10, color=GOLD, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, W - Inches(1.6), H - Inches(0.40), Inches(1.0), Inches(0.38), str(n_slides[0]), size=10, color=WHITE, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

def pic(slide, name, x, y, w=None, h=None):
    p = os.path.join(A, name + ".png")
    if not os.path.exists(p): return None
    return slide.shapes.add_picture(p, x, y, width=w, height=h)

def card(slide, x, y, w, h, title, body, fill=LIGHT, tcol=INK, tsize=15, bsize=12):
    rect(slide, x, y, w, h, fill)
    rect(slide, x, y, Inches(0.09), h, ORANGE)
    text(slide, x + Inches(0.2), y + Inches(0.1), w - Inches(0.3), Inches(0.5), title, size=tsize, bold=True, color=tcol, font=HEAD)
    text(slide, x + Inches(0.2), y + Inches(0.6), w - Inches(0.3), h - Inches(0.7), body, size=bsize, color=tcol)

# ---------------- 1. Title ----------------
s = prs.slides.add_slide(BLANK); n_slides[0] += 1
rect(s, 0, 0, W, H, BLACK); rect(s, 0, Inches(5.9), W, Inches(0.1), GOLD)
pic(s, "vguard-big-idea-logo", Inches(0.7), Inches(0.5), h=Inches(0.75))
text(s, Inches(0.7), Inches(2.0), Inches(12), Inches(1.2), "V-Guard Sentinel", size=54, bold=True, color=GOLD, font=HEAD)
text(s, Inches(0.7), Inches(3.05), Inches(12), Inches(1.0), "One AI Core that makes the home inverter and its battery self-aware — predicting failure, learning the home, acting on its own, offline.", size=22, color=WHITE)
text(s, Inches(0.7), Inches(4.4), Inches(12), Inches(0.5), "Track 4 · Reimagining V-Guard for an AI-Powered Era — From Smart Products to Intelligent Products", size=15, color=GOLD)
text(s, Inches(0.7), Inches(6.2), Inches(12), Inches(0.9), ["Team Codey Tingle (TI3271)  ·  Ayush Manoj Garg · Eshan Shukla · Tanay Chaplot", "Finale · Kochi · 24 September 2026"], size=14, color=WHITE)

# ---------------- 2. Problem ----------------
s = prs.slides.add_slide(BLANK); chrome(s, "The battery is the silent point of failure", "Every Indian home inverter is fast at switching and blind about the one part that actually fails")
pic(s, "problem_stats", Inches(0.7), Inches(1.5), w=Inches(12))
bullets(s, Inches(0.8), Inches(5.0), Inches(11.8), Inches(1.9), [
 "**The pain:** a tubular battery dies without warning — food spoils, work stops, the replacement is a ₹15–18 k surprise",
 "**The cost to V-Guard:** early failures land inside the flat warranty window; today's inverters sense voltage and current but turn none of it into foresight",
 "**The gap:** no Indian competitor ships on-device battery remaining-life prediction, autonomous load prioritisation, or offline appliance intelligence"], size=15)

# ---------------- 3. Idea ----------------
s = prs.slides.add_slide(BLANK); chrome(s, "The idea: from smart to intelligent", "Smart products take commands. Intelligent products predict, learn and act — and do it without the cloud")
pic(s, "positioning", Inches(0.7), Inches(1.5), w=Inches(12))
bullets(s, Inches(0.8), Inches(5.55), Inches(11.8), Inches(1.5), [
 "**Sentinel** is a small module — embedded in new V-Guard inverters, retrofitted to the installed base — that runs four intelligence engines on a ₹300 microcontroller",
 "**Offline-first by design:** every core function works with zero internet, which in India is a virtue, not a compromise"], size=15)

# ---------------- 4. Where it sits ----------------
s = prs.slides.add_slide(BLANK); chrome(s, "What Sentinel physically is", "A box beside the inverter: it reads the battery, watches the AC output, and switches loads through fail-safe contactors")
pic(s, "placement", Inches(0.5), Inches(1.45), w=Inches(12.3))

# ---------------- 5. Four engines ----------------
s = prs.slides.add_slide(BLANK); chrome(s, "Four engines, one core")
cw, ch = Inches(2.9), Inches(4.6); y0 = Inches(1.5)
card(s, Inches(0.6), y0, cw, ch, "1 · Battery Health", "Coulomb counting + Kalman filter fuse current, voltage and temperature into state of charge and internal resistance. A 28 k-parameter int8 network turns each outage cycle into a state-of-health band and a graded warning: “replace within 6–14 weeks”.\n\nSigned health log → one-tap, data-backed warranty.", fill=GOLD)
card(s, Inches(3.7), y0, cw, ch, "2 · Habit Autopilot", "Learns the home's load and outage rhythm in a 7×24 table. Detects a blackout in under 2 s by a 2-of-3 vote, then sheds heavy loads by tier with hysteresis so the fridge, lights and router last longest.\n\nCoil off = loads on: every fault fails safe.", fill=LIGHT)
card(s, Inches(6.8), y0, cw, ch, "3 · Energy Coach & Grid Shield", "A metering front-end on the inverter output gives real and reactive power. Event-based appliance recognition names the big loads and their kWh; half-cycle RMS logs sags, swells and interruptions as appliance-protection evidence.", fill=LIGHT)
card(s, Inches(9.9), y0, cw, ch, "4 · Fleet Intelligence", "Consented, pseudonymised per-cycle summaries (~1 KB/day) train better models centrally and return over the air. The moat is V-Guard's own fleet and aging data.\n\nFederated learning is designed (Flower, DP-FTRL) and deliberately on hold.", fill=LIGHT)
text(s, Inches(0.6), Inches(6.3), Inches(12), Inches(0.5), "Same core, same sensors, different models: the platform extends to pumps, stabilisers and water heaters.", size=13, color=GREY, italic=True)

# ---------------- 6. Engine 1 flow ----------------
s = prs.slides.add_slide(BLANK); chrome(s, "Engine 1 · How the ESP32 predicts battery life", "Physics first, learning second: the model only sees quantities the lead-acid lifetime literature says govern wear")
pic(s, "engine1_flow", Inches(0.5), Inches(1.5), w=Inches(12.3))
bullets(s, Inches(0.8), Inches(4.95), Inches(11.8), Inches(2.0), [
 "**Estimator:** 3-state extended Kalman filter (SoC, polarisation, R0); SoC re-anchors at 100 % on charge-current taper and to the OCV table only after a true rest with the charger off",
 "**Features:** resistance and sag vs the battery's own baseline, throughput, depth of discharge, coulombic efficiency, charge acceptance, dQ/dV, Arrhenius heat stress",
 "**Model:** 1-D CNN over the last 30 cycles, three seeds averaged, quantile heads P10/P50/P90 with conformal calibration — a window, never a date; 16 KB of RAM, ~1–10 ms per inference"], size=14)

# ---------------- 7. Engine 1 results ----------------
s = prs.slides.add_slide(BLANK); chrome(s, "Engine 1 · What the prototype measures today", "All numbers from synthetic tubular batteries with an independent, never-touched test split — honest, not flattering")
pic(s, "soh_band", Inches(1.9), Inches(1.4), w=Inches(9.5))
rows = [("SoH MAE / RMSE", "8.2 / 9.6 pt", "target ≤ 3 pt after the aging campaign"), ("80 % band coverage", "0.66 → 0.99 after conformal", "calibrated but too wide (34 pt) on 3 batteries"), ("EKF SoC error", "±3 % RMS", "vs 7 %/week drift for a naive counter"), ("int8 vs float", "+0.01 pt", "3 seeds → 3 × 36.7 KB .tflite")]
x0, y0 = Inches(0.9), Inches(5.3)
for i, (a, b, c) in enumerate(rows):
    y = y0 + Inches(0.38) * i
    text(s, x0, y, Inches(2.6), Inches(0.35), a, size=12, bold=True); text(s, x0 + Inches(2.6), y, Inches(3.2), Inches(0.35), b, size=12, color=ORANGE, bold=True); text(s, x0 + Inches(5.8), y, Inches(6.5), Inches(0.35), c, size=12, color=GREY)

# ---------------- 8. Engine 2 ----------------
s = prs.slides.add_slide(BLANK); chrome(s, "Engine 2 · Autopilot: essentials alive, fail-safe always")
pic(s, "autopilot_ladder", Inches(0.5), Inches(1.4), w=Inches(7.4))
pic(s, "failsafe_flow", Inches(0.5), Inches(4.45), w=Inches(7.4))
bullets(s, Inches(8.4), Inches(1.5), Inches(4.6), Inches(5.3), [
 "**Learns, transparently:** hour-of-week load and outage tables (4.7 KB) — a judge can read them; single-home load is too noisy for a neural forecaster to beat them",
 "**Forecasts the roster, not the storm:** scheduled cuts are predictable from the home's own history; faults are detected, not predicted",
 "**Tiers:** T1 never shed · T2 defer at 55 % · T3 shed at 40 %, 15-point hysteresis, 3–5 min compressor dwell, 30-min user override that cannot beat the safety floor",
 "**Measured:** 65 tests, 50/50 scripted scenarios in C, zero chatter, zero dwell violations"], size=13)

# ---------------- 9. Engine 3 ----------------
s = prs.slides.add_slide(BLANK); chrome(s, "Engine 3 · Energy Coach and Grid Shield", "Both need one thing the report's first BOM lacked: a real metering front-end on the AC output — now designed in")
pic(s, "nilm_events", Inches(0.5), Inches(1.45), w=Inches(6.2))
pic(s, "pq_dip", Inches(6.8), Inches(1.45), w=Inches(6.2))
bullets(s, Inches(0.7), Inches(4.05), Inches(6.0), Inches(3.0), [
 "**How the ESP32 tells appliances apart:** each ON/OFF step becomes a 13-number signature — ΔP, ΔQ, phase angle, inrush ratio, settling time, harmonic share, duration, periodicity",
 "**Per-home library:** rules + k-NN, seeded from an Indian appliance table, grown by clustering, named by the user (“what just turned on?”)",
 "**Honest scope:** fridge, AC, geyser, pump, iron/kettle class — not two identical fans, not loads under 40 W"], size=12)
bullets(s, Inches(7.0), Inches(4.05), Inches(6.0), Inches(3.0), [
 "**Half-cycle RMS** on an isolated tap, IEEE 1159 buckets, 4-cycle FFT for THD; Class-S-like, never “Class A”",
 "**Measured:** 13/13 dips detected, magnitude error < 1 %, duration error 0 half-cycles; event recall 1.00, fridge F1 0.97",
 "**Value:** a dated, signed record of utility-caused events — appliance-protection evidence, not a trip"], size=12)

# ---------------- 10. Hardware & product ----------------
s = prs.slides.add_slide(BLANK); chrome(s, "The product: two SKUs, one board", "Embedded in new V-Guard inverters, or a 30-minute electrician retrofit on the installed base")
pic(s, "cost_tiers", Inches(0.5), Inches(1.45), w=Inches(7.3))
card(s, Inches(8.1), Inches(1.5), Inches(4.8), Inches(2.45), "Sentinel-Embedded", "Daughter-board inside the inverter. Kelvin-taps the inverter's own shunt; controls the charger through a DAC on the PWM feedback node with a hardware voltage/thermal ceiling. Adaptive charging and the life-extension claim live here only.", fill=GOLD, bsize=11.5)
card(s, Inches(8.1), Inches(4.15), Inches(4.8), Inches(2.45), "Sentinel-Retrofit", "Box beside the inverter: external 0.1 mΩ shunt in the battery negative lead, NTC on the battery, CT + voltage tap on AC-OUT, 4 fail-safe contactor channels. Advisory on charging, protective on loads. ~2.6 mA average draw — 0.04 % of a 150 Ah battery per day.", fill=LIGHT, bsize=11.5)
text(s, Inches(0.7), Inches(5.9), Inches(7.0), Inches(1.0), "Parts added versus the submitted report: ATM90E32AS metering AFE + CT, DS3231 clock, ATECC608 secure element, driver-board supervisory timer. Every part is tied to the engine it enables.", size=12, color=GREY, italic=True)

# ---------------- 11. Prototype built ----------------
s = prs.slides.add_slide(BLANK); chrome(s, "The prototype: what actually runs", "Not a slide deck — a codebase: every engine implemented in Python and C, tested, with the same C linked into the firmware")
pic(s, "tests_per_module", Inches(0.5), Inches(1.45), w=Inches(7.0))
bullets(s, Inches(7.8), Inches(1.5), Inches(5.2), Inches(5.4), [
 "**Simulator + features:** synthetic Indian tubular duty cycles at 1 Hz → 14 + 6 physics features per outage cycle",
 "**Model pipeline:** train → evaluate → int8 quantise → full-int8 .tflite (3 seeds), independent test split, leakage guard",
 "**EKF, autopilot, NILM, PQ, health log, charger policy:** Python reference + C99 port, bit-identical outputs",
 "**Firmware:** ESP-IDF project with the real task map; a host build links every C module, passes an int8 golden self-test and feeds the dashboard",
 "**Dashboard:** Flask + one page; scripted finale sequence or live file from the device",
 "**Reproducible:** run_all.sh regenerates every number; pinned dependencies; private GitHub repo"], size=13)

# ---------------- 12. Self-contained ----------------
s = prs.slides.add_slide(BLANK); chrome(s, "Self-contained by construction", "Unplug the router for a year: every row below still works. The cloud only adds updates and fleet learning")
pic(s, "firmware_tasks", Inches(0.5), Inches(1.45), w=Inches(12.3))
text(s, Inches(0.7), Inches(6.35), Inches(12), Inches(0.6), "Fixed int8 weights in flash · 16 KB arena · one inference per outage cycle · plain C turns quantiles into “replace in 6–14 weeks” · no training on the device, ever.", size=13, color=GREY, italic=True)

# ---------------- 13. Demo ----------------
s = prs.slides.add_slide(BLANK); chrome(s, "Live demonstration", "A real 12 V tubular battery and inverter on the table; three relay-switched load tiers; nothing connected to the internet")
pic(s, "demo_flow", Inches(0.4), Inches(1.6), w=Inches(12.5))
bullets(s, Inches(0.8), Inches(4.6), Inches(11.8), Inches(2.3), [
 "**Real:** SoC and internal resistance from the shunt and Kalman filter; outage detection; tiered shedding; fail-safe restore on reset; appliance events; signed log verification",
 "**Replayed and labelled as such:** the state-of-health walk to “Replace within N weeks” uses a synthetic aging trajectory — a bench battery does not age in a week",
 "**Deliberately absent:** federated learning (on hold), charger control on a retrofit inverter (no interface exists), loads under 40 W"], size=14)

# ---------------- 14. Evidence ladder ----------------
s = prs.slides.add_slide(BLANK); chrome(s, "Proven · synthetic · still needs hardware", "We separate what is deployable today from what is a research target — credibility is the whole game")
pic(s, "evidence_ladder", Inches(0.5), Inches(1.5), w=Inches(12.3))
bullets(s, Inches(0.8), Inches(5.6), Inches(11.8), Inches(1.4), [
 "**Corrections we made to our own report after building it:** float voltage is not open-circuit voltage; a battery shunt cannot see appliances; adaptive charging is embedded-only; federate the model that has labels on the device",
 "**Every claim traces to a script** in the repository and to a design decision (D1–D17) in the design register"], size=13)

# ---------------- 15. Roadmap ----------------
s = prs.slides.add_slide(BLANK); chrome(s, "Roadmap as validation gates, and the ask")
gates = [("Gate 1 · Dataset", "24–36 tubular batteries, 3 temperatures × 3 depths × 2 regimes, 6–8 months in the Kochi reliability lab. Output: EKF tables, real SoH labels, the measured life-extension ratio.", GOLD),
         ("Gate 2 · Field pilot", "100 retrofit units, 6 months: SoC within 3 % on reference shunts, ≥ 75 % RUL-window hit-rate, zero unsafe sheds, fridge/AC F1 ≥ 0.8, power within budget.", LIGHT),
         ("Gate 3 · Learning stability", "Three model releases with no regression on a golden slice; embedded SKU into the Smart Pro family with its existing Wi-Fi/BLE and app.", LIGHT)]
for i, (t, b, c) in enumerate(gates):
    card(s, Inches(0.6) + Inches(4.15) * i, Inches(1.5), Inches(3.95), Inches(3.3), t, b, fill=c, bsize=12)
bullets(s, Inches(0.8), Inches(5.1), Inches(11.8), Inches(1.8), [
 "**What V-Guard gains:** timely battery replacement revenue, warranty decided from evidence, a premium “intelligent” tier, and a fleet-data moat no competitor can buy",
 "**What we need:** the ESP32-S3 bench hardware, one opened Prime-series inverter with its charger-board schematic, batteries and cyclers for Gate 1"], size=14)

# ---------------- 16. Close ----------------
s = prs.slides.add_slide(BLANK); n_slides[0] += 1
rect(s, 0, 0, W, H, BLACK); rect(s, 0, Inches(3.55), W, Inches(0.1), GOLD)
text(s, Inches(0.7), Inches(1.6), Inches(12), Inches(1.4), "Coil off means loads on.\nA window, never a date.\nOffline first, fleet second.", size=34, bold=True, color=GOLD, font=HEAD)
text(s, Inches(0.7), Inches(4.0), Inches(12), Inches(1.0), "V-Guard Sentinel — Team Codey Tingle (TI3271)", size=20, color=WHITE)
text(s, Inches(0.7), Inches(4.7), Inches(12), Inches(1.6), ["Design register: 17 decisions · Prototype: 245 automated tests, 6 C modules, full-int8 model, firmware host build", "github.com/epshukla/vguard-sentinel (private)"], size=14, color=GREY)
pic(s, "vguard-big-idea-logo", Inches(0.7), Inches(6.4), h=Inches(0.6))

prs.save("VGuard-Sentinel-Finale-Deck.pptx"); print("saved, slides:", len(prs.slides))
